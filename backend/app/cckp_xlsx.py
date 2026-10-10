"""Read the inspected CCKP export format. Ambiguous time columns fail closed.

Filename product metadata describes the requested product, not proof that a date-style
column represents that entire period. Such candidates remain outside the repository.
No network calls or implicit geography matching are performed.
"""
from datetime import datetime, timezone
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
import json
import math
from pathlib import Path
import re

from openpyxl import load_workbook

from .providers import validate_projection
from .projections import ROOT, store_projection_records

DOWNLOAD_URL = 'https://climateknowledgeportal.worldbank.org/download-data'
METADATA_URL = 'https://climateknowledgeportal.worldbank.org/index.php/metadata'
COLLECTION_URL = 'https://worldbank.github.io/climateknowledgeportal/docs/collections/cmip6-x0.25.html'
ADAPTER_VERSION = 'cckp-xlsx-1'
PRODUCT = re.compile(
    r'^cmip6-x0\.25_climatology_(?P<variable>[^_]+)_climatology_annual_'
    r'(?P<periods>1995-2014(?:,2040-2059)?)_median_'
    r'(?P<scenario>historical|ssp245)_ensemble_all_mean\.xlsx$')
GEOGRAPHIES = {'IND': 'India'}  # Provider country example explicitly identifies IND.


def content_hash(value):
    return sha256(json.dumps(value, sort_keys=True, separators=(',', ':'),
                             ensure_ascii=False, allow_nan=False).encode()).hexdigest()


def inspect_cckp_workbook(path):
    """Return validated schema candidates and explicit semantic verification status."""
    path = Path(path)
    match = PRODUCT.fullmatch(path.name)
    if not match:
        raise ValueError('Unsupported CCKP filename/product; requires a separately inspected adapter')
    if match['variable'] != 'tasmax':
        raise ValueError('Unsupported variable: this adapter supports tasmax only')
    scenario = match['scenario']
    period = '1995-2014' if scenario == 'historical' else '2040-2059'
    if match['periods'] != ('1995-2014' if scenario == 'historical' else '1995-2014,2040-2059'):
        raise ValueError('Scenario/filename period mismatch')
    raw = path.read_bytes()
    digest = sha256(raw).hexdigest()
    try:
        book = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise ValueError('Malformed XLSX workbook') from exc
    try:
        expected_sheet = 'all' if scenario == 'historical' else period
        if book.sheetnames != [expected_sheet]:
            raise ValueError(f'Expected only worksheet {expected_sheet!r}; found {book.sheetnames!r}')
        sheet = book[expected_sheet]
        if sheet.max_row > 10000 or sheet.max_column != 3:
            raise ValueError('Unsupported worksheet dimensions')
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < 2 or tuple(rows[0][:2]) != ('code', 'name'):
            raise ValueError('Missing code/name header or data rows')
        label = rows[0][2]
        if not isinstance(label, str):
            raise ValueError('Time header must preserve an explicit source string')
        if label not in (period, period[:4] + '-07'):
            raise ValueError('Unsupported or conflicting time-column label')
        verified = label == period
        reason = None if verified else (
            f'Time column {label!r} has no verified mapping to annual period {period}. '
            'CCKP documentation defines annual climatology but does not document this export date encoding.')
        records = []
        seen = set()
        for index, row in enumerate(rows[1:], 2):
            if all(v is None for v in row):
                continue
            code, name, value = row
            if code not in GEOGRAPHIES or GEOGRAPHIES[code] != name:
                raise ValueError(f'Unknown or inconsistent geography at row {index}: {code!r}/{name!r}')
            if code in seen:
                raise ValueError('Duplicate geography records')
            seen.add(code)
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
                raise ValueError(f'Missing, nonnumeric or nonfinite value at row {index}')
            record = validate_projection({
                'regionId': code, 'geographyName': name, 'geoLevel': 'national',
                'collection': 'cmip6-x0.25', 'variable': 'tasmax', 'unit': '°C',
                'scenario': scenario, 'period': period, 'product': 'annual climatology',
                'ensemble': 'multi-model', 'percentile': 'p50', 'value': value,
                'baseline': '1995-2014', 'sourceURL': DOWNLOAD_URL, 'downloadedAt': None,
                'sourceFile': path.name, 'sourceFileHash': digest,
                'rawColumnLabels': list(rows[0]), 'sourceSheet': expected_sheet,
                'sourceCell': f'C{index}', 'sourceModel': 'ensemble',
                'sourceModelCalculation': 'all', 'sourceStatistic': 'mean',
                'sourcePercentile': 'median', 'metadataURLs': [METADATA_URL, COLLECTION_URL, DOWNLOAD_URL],
                'metadataBasis': 'Requested product filename and official CCKP product documentation',
                'verificationStatus': 'verified' if verified else 'awaiting_time_label_verification',
                'verificationReason': reason, 'normalizationVersion': ADAPTER_VERSION,
                'companyLocationVerified': False, 'usedInSimulation': False,
            })
            record['normalizedContentHash'] = content_hash(record)
            records.append(record)
        if not records:
            raise ValueError('No projection data rows')
        # Excel document creation time is not evidence of an actual download time.
        properties = {key: str(getattr(book.properties, key)) if getattr(book.properties, key) is not None else None
                      for key in ('creator', 'title', 'description', 'created', 'modified')}
        result = {
            'schemaVersion': '1.0.0', 'adapterVersion': ADAPTER_VERSION,
            'sha256': digest, 'sourceFile': path.name, 'sourceSheet': expected_sheet,
            'rawColumnLabels': list(rows[0]), 'workbookProperties': properties,
            'verificationStatus': 'verified' if verified else 'awaiting_time_label_verification',
            'reason': reason, 'candidateRecords': records,
            'normalizedContentHash': content_hash(records),
            'importedAt': datetime.now(timezone.utc).isoformat(), 'downloadedAt': None,
        }
    finally:
        book.close()
    if path.read_bytes() != raw:
        raise ValueError('Source changed during inspection')
    return result


def compare_context(historical, future):
    """Subtract comparable verified period medians, never a model anomaly or loss."""
    for record in (historical, future):
        validate_projection(record)
        if record.get('verificationStatus') != 'verified':
            return {'status': 'unavailable', 'difference': None,
                    'reason': 'Annual time-column interpretation is awaiting verification'}
    fields = ('regionId', 'geoLevel', 'collection', 'variable', 'unit', 'product',
              'ensemble', 'percentile', 'baseline', 'sourceStatistic', 'sourceModelCalculation')
    if any(historical.get(key) != future.get(key) for key in fields):
        raise ValueError('Historical/future context is not comparable')
    if (historical['scenario'], historical['period'], future['scenario'], future['period']) != (
            'historical', '1995-2014', 'ssp245', '2040-2059'):
        raise ValueError('Historical/future comparison periods or scenarios mismatch')
    if any(record['value'] is None for record in (historical, future)):
        return {'status': 'unavailable', 'difference': None, 'reason': 'Missing value'}
    return {'status': 'available', 'difference': float(Decimal(str(future['value'])) - Decimal(str(historical['value']))),
            'unit': historical['unit'], 'historicalHash': historical['normalizedContentHash'],
            'futureHash': future['normalizedContentHash'],
            'limitation': 'Difference of period ensemble medians, not median model anomaly, financial loss or uncertainty interval'}


def inspect_pair(historical, future):
    reports = [inspect_cckp_workbook(historical), inspect_cckp_workbook(future)]
    if [r['candidateRecords'][0]['scenario'] for r in reports] != ['historical', 'ssp245']:
        raise ValueError('Historical/future inputs are reversed or inconsistent')
    codes = [{r['regionId'] for r in report['candidateRecords']} for report in reports]
    if codes[0] != codes[1]:
        raise ValueError('Historical/future geography sets mismatch')
    comparisons = [compare_context(h, next(f for f in reports[1]['candidateRecords'] if f['regionId'] == h['regionId']))
                   for h in reports[0]['candidateRecords']]
    return {'sources': reports, 'comparisons': comparisons,
            'status': 'verified' if all(r['verificationStatus'] == 'verified' for r in reports) else 'awaiting_verification'}


def install_pair(historical, future, root=ROOT):
    report = inspect_pair(historical, future)
    if report['status'] != 'verified':
        raise ValueError('Installation blocked: unresolved time labels; use --quarantine to retain the audit')
    mapping = {'adapter': ADAPTER_VERSION, 'geographyMap': GEOGRAPHIES}
    # Validate both sources against existing products before writing either one.
    for path, source in zip((historical, future), report['sources']):
        store_projection_records(path, mapping, source['candidateRecords'], root, validate_only=True,
                                 metadata={'normalizedContentHash': source['normalizedContentHash'],
                                           'importedAt': source['importedAt']})
    for path, source in zip((historical, future), report['sources']):
        store_projection_records(path, mapping, source['candidateRecords'], root,
                                 metadata={'normalizedContentHash': source['normalizedContentHash'],
                                           'importedAt': source['importedAt']})
    return report


def quarantine_pair(historical, future, output):
    """Store only under a nested quarantine folder, never a queryable root manifest."""
    report = inspect_pair(historical, future)
    target = Path(output) / 'quarantine'
    digest = content_hash([s['normalizedContentHash'] for s in report['sources']])
    destination = target / (digest + '.json')
    # Check all conflicts before writes; repeat audits preserve the first import time.
    raw_files = []
    for path, source in zip((historical, future), report['sources']):
        raw = Path(path).read_bytes()
        if sha256(raw).hexdigest() != source['sha256']:
            raise ValueError('Source changed after inspection')
        copy = target / 'raw' / (source['sha256'] + '.xlsx')
        if copy.exists() and copy.read_bytes() != raw:
            raise ValueError('Quarantine raw integrity conflict')
        raw_files.append((copy, raw))
    if destination.exists():
        previous = json.loads(destination.read_text(encoding='utf-8'))
        if [s['normalizedContentHash'] for s in previous['sources']] != [s['normalizedContentHash'] for s in report['sources']]:
            raise ValueError('Quarantine audit integrity conflict')
        if any(not copy.exists() for copy, _ in raw_files):
            raise ValueError('Quarantine raw copy missing')
        return previous, destination
    (target / 'raw').mkdir(parents=True, exist_ok=True)
    for copy, raw in raw_files:
        if not copy.exists():
            copy.write_bytes(raw)
    with destination.open('x', encoding='utf-8') as stream:
        json.dump(report, stream, indent=2, ensure_ascii=False, allow_nan=False)
    return report, destination
