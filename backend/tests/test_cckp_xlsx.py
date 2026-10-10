"""Actual source workbooks plus disposable structural mutations, never installed fake data."""
from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
import subprocess
import sys

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from app.cckp_xlsx import (inspect_cckp_workbook, inspect_pair, compare_context,
                           install_pair, quarantine_pair, content_hash)
from app.main import app, repository
from app.projections import ProjectionRepository

PROJECT = Path(__file__).parents[1]
SOURCES = PROJECT / 'data/incoming/cckp'
HIST = SOURCES / 'cmip6-x0.25_climatology_tasmax_climatology_annual_1995-2014_median_historical_ensemble_all_mean.xlsx'
FUTURE = SOURCES / 'cmip6-x0.25_climatology_tasmax_climatology_annual_1995-2014,2040-2059_median_ssp245_ensemble_all_mean.xlsx'
HASHES = ['df46b53cdc00dedfd080f216c07d27f028f83b75638231ba767b4ed9cadb4999',
          '0bb10cdca67f933c022a76af48f03e07e1f845112ff9535c650f9053c154c5ac']


def mutate(source, folder, change):
    """Only mutate temporary copies to exercise rejection/accepted schema branches."""
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / source.name
    book = load_workbook(source)
    try:
        change(book)
        book.save(target)
    finally:
        book.close()
    return target


@pytest.mark.parametrize('path,sheet,label,value,scenario,period,digest', [
    (HIST, 'all', '1995-07', 29.68, 'historical', '1995-2014', HASHES[0]),
    (FUTURE, '2040-2059', '2040-07', 30.77, 'ssp245', '2040-2059', HASHES[1]),
])
def test_genuine_workbook_schema_values_provenance(path, sheet, label, value, scenario, period, digest):
    original = path.read_bytes()
    report = inspect_cckp_workbook(path)
    assert report['sourceSheet'] == sheet
    assert report['rawColumnLabels'] == ['code', 'name', label]
    assert report['sha256'] == digest == sha256(original).hexdigest()
    assert len(report['candidateRecords']) == 1
    r = report['candidateRecords'][0]
    assert (r['regionId'], r['geographyName'], r['geoLevel']) == ('IND', 'India', 'national')
    assert (r['value'], r['variable'], r['unit'], r['scenario'], r['period']) == (value, 'tasmax', '°C', scenario, period)
    assert (r['ensemble'], r['percentile'], r['sourceStatistic']) == ('multi-model', 'p50', 'mean')
    assert r['downloadedAt'] is None and report['downloadedAt'] is None
    assert r['sourceFileHash'] == digest and r['sourceCell'] == 'C2'
    assert r['metadataURLs'] and all(url.startswith('https://') for url in r['metadataURLs'])
    assert not r['companyLocationVerified'] and not r['usedInSimulation']
    assert report['importedAt'] and r['verificationStatus'] == 'awaiting_time_label_verification'
    assert r['normalizedContentHash'] == content_hash({k: v for k, v in r.items() if k != 'normalizedContentHash'})
    assert path.read_bytes() == original


def test_date_encoding_does_not_become_july_observation_or_verified_delta():
    report = inspect_pair(HIST, FUTURE)
    assert report['status'] == 'awaiting_verification'
    assert report['comparisons'][0]['status'] == 'unavailable'
    assert report['comparisons'][0]['difference'] is None
    assert all('no verified mapping' in s['reason'] for s in report['sources'])


def test_normalization_is_deterministic_and_preserves_unknown_download_time():
    a, b = inspect_pair(HIST, FUTURE), inspect_pair(HIST, FUTURE)
    assert [s['normalizedContentHash'] for s in a['sources']] == [s['normalizedContentHash'] for s in b['sources']]
    assert [s['candidateRecords'] for s in a['sources']] == [s['candidateRecords'] for s in b['sources']]


def test_unverified_pair_cannot_install_or_create_root(tmp_path):
    root = tmp_path / 'active'
    with pytest.raises(ValueError, match='Installation blocked'):
        install_pair(HIST, FUTURE, root)
    assert not root.exists()


def test_quarantine_duplicate_raw_integrity_and_repository_exclusion(tmp_path):
    before = [p.read_bytes() for p in (HIST, FUTURE)]
    report, path = quarantine_pair(HIST, FUTURE, tmp_path)
    saved = path.read_bytes()
    duplicate, same_path = quarantine_pair(HIST, FUTURE, tmp_path)
    assert same_path == path and duplicate == report and path.read_bytes() == saved
    assert len(list(path.parent.glob('*.json'))) == 1
    for source, raw in zip(report['sources'], before):
        assert (path.parent / 'raw' / (source['sha256'] + '.xlsx')).read_bytes() == raw
    assert ProjectionRepository(tmp_path).query('IND', 'tasmax', '2040-2059', 'ssp245') == []
    assert [p.read_bytes() for p in (HIST, FUTURE)] == before


def test_quarantine_raw_conflict_is_rejected(tmp_path):
    report, path = quarantine_pair(HIST, FUTURE, tmp_path)
    raw = path.parent / 'raw' / (report['sources'][0]['sha256'] + '.xlsx')
    raw.write_bytes(b'corrupt')
    with pytest.raises(ValueError, match='integrity conflict'):
        quarantine_pair(HIST, FUTURE, tmp_path)


@pytest.mark.parametrize('change,message', [
    (lambda b: setattr(b.active, 'title', 'wrong'), 'Expected only worksheet'),
    (lambda b: b.create_sheet('extra'), 'Expected only worksheet'),
    (lambda b: b.active.__setitem__('A1', 'unknown'), 'Missing code/name'),
    (lambda b: b.active.__setitem__('C1', '2050-07'), 'conflicting time'),
    (lambda b: b.active.__setitem__('A2', 'UNKNOWN'), 'Unknown or inconsistent geography'),
    (lambda b: b.active.__setitem__('B2', 'Incorrect country'), 'Unknown or inconsistent geography'),
    (lambda b: b.active.__setitem__('C2', None), 'Missing, nonnumeric'),
    (lambda b: b.active.__setitem__('C2', '=1+1'), 'Missing, nonnumeric'),
    (lambda b: b.active.__setitem__('C2', True), 'Missing, nonnumeric'),
    (lambda b: b.active.append(['IND', 'India', 29.68]), 'Duplicate geography'),
    (lambda b: b.active.delete_rows(2), 'Missing code/name'),
])
def test_rejects_ambiguous_or_malformed_structure(tmp_path, change, message):
    source = mutate(HIST, tmp_path, change)
    with pytest.raises(ValueError, match=message):
        inspect_cckp_workbook(source)


def test_malformed_zip_and_unsupported_variable(tmp_path):
    invalid = tmp_path / HIST.name
    invalid.write_bytes(b'not excel')
    with pytest.raises(ValueError, match='Malformed XLSX'):
        inspect_cckp_workbook(invalid)
    unsupported = tmp_path / HIST.name.replace('tasmax', 'pr')
    unsupported.write_bytes(HIST.read_bytes())
    with pytest.raises(ValueError, match='Unsupported variable'):
        inspect_cckp_workbook(unsupported)


def explicit_period_copies(tmp_path):
    # Accepted schema branch only: genuine values with unambiguous full-period headers
    # in disposable test copies. These are not provider evidence or production sources.
    h = mutate(HIST, tmp_path, lambda b: b.active.__setitem__('C1', '1995-2014'))
    f = mutate(FUTURE, tmp_path, lambda b: b.active.__setitem__('C1', '2040-2059'))
    return h, f


def test_verified_schema_comparison_and_install_idempotency(tmp_path):
    h, f = explicit_period_copies(tmp_path / 'validation-fixtures')
    root = tmp_path / 'installed'
    result = inspect_pair(h, f)
    assert result['comparisons'][0]['difference'] == 1.09
    install_pair(h, f, root)
    before = {p.name: p.read_bytes() for p in root.glob('*.json')}
    install_pair(h, f, root)
    assert {p.name: p.read_bytes() for p in root.glob('*.json')} == before
    assert len(list(root.glob('*.json'))) == 2
    assert ProjectionRepository(root).query('IND', 'tasmax', '1995-2014', 'historical')[0]['value'] == 29.68
    assert ProjectionRepository(root).query('IND', 'tasmax', '2040-2059', 'ssp245')[0]['value'] == 30.77
    assert ProjectionRepository(root).query('IND', 'hd35', '2040-2059', 'ssp245') == []
    assert ProjectionRepository(root).query('IND', 'tasmax', '2040-2059', 'ssp585') == []
    assert ProjectionRepository(root).query('UNKNOWN', 'tasmax', '2040-2059', 'ssp245') == []


@pytest.mark.parametrize('key,value', [('regionId', 'OTHER'), ('unit', 'K'), ('percentile', 'p90'),
                                     ('sourceStatistic', 'max'), ('baseline', '1986-2005'), ('period', '1995-2014')])
def test_comparison_requires_matching_products(tmp_path, key, value):
    h, f = explicit_period_copies(tmp_path)
    report = inspect_pair(h, f)
    historical = report['sources'][0]['candidateRecords'][0]
    future = {**report['sources'][1]['candidateRecords'][0], key: value}
    with pytest.raises(ValueError):
        compare_context(historical, future)


def test_conflicting_install_does_not_partially_install(tmp_path):
    h, f = explicit_period_copies(tmp_path / 'validation-fixtures')
    root = tmp_path / 'installed'
    install_pair(h, f, root)
    before = {p.name: p.read_bytes() for p in root.glob('*.json')}
    changed = mutate(f, tmp_path / 'changed', lambda b: b.active.__setitem__('B2', 'India'))
    # Changing an unrelated property still creates a different immutable raw source.
    book = load_workbook(changed)
    book.properties.title = 'Different source'
    book.save(changed)
    book.close()
    with pytest.raises(ValueError, match='already imported'):
        install_pair(h, changed, root)
    assert {p.name: p.read_bytes() for p in root.glob('*.json')} == before


def test_real_sources_remain_unavailable_via_api_without_synthetic_promotion(tmp_path, monkeypatch):
    quarantine_pair(HIST, FUTURE, tmp_path)
    monkeypatch.setattr('app.climate_api.projections', ProjectionRepository(tmp_path))
    datasets_before = [deepcopy(repository.dataset(key)) for key in ('reference', 'climateweb')]
    with TestClient(app) as client:
        for key, data in zip(('reference', 'climateweb'), datasets_before):
            for period, scenario in [('1995-2014', 'historical'), ('2040-2059', 'ssp245')]:
                response = client.get('/api/v1/projections', params={'datasetId': key, 'regionId': data['regions'][0]['id'],
                    'variable': 'tasmax', 'period': period, 'scenario': scenario})
                assert response.status_code == 200
                assert response.json()['status'] == 'unavailable' and response.json()['records'] == []
                assert 'No verified provider-geography mapping' in response.json()['reason']
            assert client.get('/api/v1/projections', params={'datasetId': key, 'regionId': 'IND'}).status_code == 422
    assert [repository.dataset(key) for key in ('reference', 'climateweb')] == datasets_before


def test_cli_dry_run_no_writes_and_install_failure(tmp_path):
    command = [sys.executable, '-m', 'scripts.import_cckp_xlsx', '--historical', str(HIST),
               '--future', str(FUTURE), '--output', str(tmp_path / 'active')]
    dry = subprocess.run(command, cwd=PROJECT, capture_output=True, text=True)
    assert dry.returncode == 0 and json.loads(dry.stdout)['action'] == 'validate-only'
    assert not (tmp_path / 'active').exists()
    install = subprocess.run(command + ['--install'], cwd=PROJECT, capture_output=True, text=True)
    assert install.returncode == 2 and 'Installation blocked' in install.stderr
    assert not (tmp_path / 'active').exists()


def test_cli_missing_inputs_clear_error(tmp_path):
    result = subprocess.run([sys.executable, '-m', 'scripts.import_cckp_xlsx', '--historical', str(tmp_path / HIST.name),
        '--future', str(FUTURE)], cwd=PROJECT, capture_output=True, text=True)
    assert result.returncode == 2 and 'CCKP import failed' in result.stderr
